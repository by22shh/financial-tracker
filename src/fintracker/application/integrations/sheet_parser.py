"""Разбор исходной таблицы Google Sheets (FR-63, FR-64, SOURCE_ANALYSIS).

Дневная матрица идёт от 10-го числа к 9-му следующего месяца. Одна ненулевая
дневная ячейка становится дневным агрегатом; итоговые столбцы и строки итогов
не импортируются вторым набором расходов.
"""

from __future__ import annotations

import datetime as dt
import hashlib
import re
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from fintracker.core.errors import ValidationFailed

# Заголовок листа вида «10.08 - 09.09» задаёт границы периода.
_SHEET_TITLE = re.compile(
    r"^\s*(?P<d1>\d{1,2})[.](?P<m1>\d{1,2})\s*[-–—]\s*(?P<d2>\d{1,2})[.](?P<m2>\d{1,2})\s*$"
)
# Ячейка может содержать выражение: несколько слагаемых или деление доли.
_EXPRESSION = re.compile(r"^=?\s*[\d\s+\-*/().,]+$")


@dataclass(frozen=True, slots=True)
class SheetDailyCell:
    """Ненулевая дневная ячейка — один будущий агрегат."""

    sheet_name: str
    row_label: str
    row_index: int
    column_letter: str
    day_offset: int
    occurred_date: dt.date
    amount: Decimal
    formula: str | None


@dataclass(frozen=True, slots=True)
class SheetRow:
    """Бюджетная строка листа с планом и фактом."""

    index: int
    label: str
    group: str | None
    detail: str | None
    fact: Decimal | None
    plan: Decimal | None
    deviation: Decimal | None


@dataclass(frozen=True, slots=True)
class SheetIncome:
    label: str
    amount: Decimal
    occurred_date: dt.date | None


@dataclass(frozen=True, slots=True)
class ParsedSheet:
    name: str
    period_start: dt.date
    period_end_exclusive: dt.date
    rows: tuple[SheetRow, ...]
    daily_cells: tuple[SheetDailyCell, ...]
    incomes: tuple[SheetIncome, ...]
    totals: dict[str, Decimal] = field(default_factory=dict)

    @property
    def nonzero_cells(self) -> int:
        return len(self.daily_cells)

    def daily_total(self) -> Decimal:
        return sum((cell.amount for cell in self.daily_cells), Decimal(0))

    def rows_total(self) -> Decimal:
        return sum((row.fact or Decimal(0) for row in self.rows), Decimal(0))


@dataclass(frozen=True, slots=True)
class ParsedWorkbook:
    sheets: tuple[ParsedSheet, ...]
    snapshot_hash: str
    taken_at: dt.datetime
    source_name: str


def _to_decimal(value: Any) -> Decimal | None:
    """Точное десятичное значение; float не используется в денежном пути."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, int):
        return Decimal(value)
    if isinstance(value, float):
        # openpyxl возвращает float; преобразование идёт через строку, чтобы
        # не тянуть двоичную погрешность дальше (FR-26).
        return Decimal(repr(value))
    if isinstance(value, Decimal):
        return value
    text = str(value).strip().replace(" ", "").replace(" ", "").replace(",", ".")
    if not text:
        return None
    try:
        return Decimal(text)
    except InvalidOperation:
        return None


def parse_sheet_title(title: str, *, year: int) -> tuple[dt.date, dt.date] | None:
    """Границы периода из названия листа «10.08 - 09.09»."""
    match = _SHEET_TITLE.match(title)
    if match is None:
        return None
    start = dt.date(year, int(match.group("m1")), int(match.group("d1")))
    end_inclusive_month = int(match.group("m2"))
    end_year = year + 1 if end_inclusive_month < start.month else year
    end_inclusive = dt.date(end_year, end_inclusive_month, int(match.group("d2")))
    return start, end_inclusive + dt.timedelta(days=1)


def parse_workbook(path: Path, *, year: int, taken_at: dt.datetime) -> ParsedWorkbook:
    """Разобрать XLSX без исполнения формул (SOURCE_ANALYSIS §5)."""
    from openpyxl import load_workbook

    raw = path.read_bytes()
    snapshot_hash = hashlib.sha256(raw).hexdigest()
    values = load_workbook(path, data_only=True, read_only=False)
    formulas = load_workbook(path, data_only=False, read_only=False)

    sheets: list[ParsedSheet] = []
    for sheet_name in values.sheetnames:
        bounds = parse_sheet_title(sheet_name, year=year)
        if bounds is None:
            continue
        sheets.append(
            _parse_sheet(values[sheet_name], formulas[sheet_name], name=sheet_name, bounds=bounds)
        )
    if not sheets:
        raise ValidationFailed(
            "В файле не найдено листов с расчётным периодом вида «10.08 - 09.09»"
        )
    return ParsedWorkbook(
        sheets=tuple(sheets),
        snapshot_hash=snapshot_hash,
        taken_at=taken_at,
        source_name=path.name,
    )


def _merged_group_labels(sheet: Any, column: int) -> dict[int, str]:
    """Заголовок группы протягивается только в пределах объединения (FR-63).

    Разные группы «Другое» не объединяются по одному слову.
    """
    labels: dict[int, str] = {}
    for merged in sheet.merged_cells.ranges:
        if merged.min_col > column or merged.max_col < column:
            continue
        value = sheet.cell(row=merged.min_row, column=column).value
        text = str(value or "").strip()
        if not text:
            continue
        for row_index in range(merged.min_row, merged.max_row + 1):
            labels[row_index] = text
    return labels


def _parse_sheet(
    values_sheet: Any, formula_sheet: Any, *, name: str, bounds: tuple[dt.date, dt.date]
) -> ParsedSheet:
    start, end_exclusive = bounds
    header_row = _find_day_header_row(values_sheet)
    if header_row is None:
        raise ValidationFailed(f"На листе «{name}» не найдена дневная матрица")

    day_columns = _day_columns(values_sheet, header_row)
    group_labels = _merged_group_labels(values_sheet, column=3)
    rows: list[SheetRow] = []
    cells: list[SheetDailyCell] = []
    incomes: list[SheetIncome] = []

    for row_index in range(header_row + 1, values_sheet.max_row + 1):
        own_group = str(values_sheet.cell(row=row_index, column=3).value or "").strip()
        group = own_group or group_labels.get(row_index, "")
        detail = str(values_sheet.cell(row=row_index, column=5).value or "").strip()
        label = " / ".join(part for part in (group, detail) if part)
        fact = _to_decimal(values_sheet.cell(row=row_index, column=6).value)
        plan = _to_decimal(values_sheet.cell(row=row_index, column=7).value)
        deviation = _to_decimal(values_sheet.cell(row=row_index, column=8).value)

        # Строка без подписи не является бюджетной: так выглядят итоговые
        # строки F/H/J, которые не импортируются вторым набором (FR-64, A74).
        if not label:
            continue
        if _looks_like_total_row(label):
            continue

        rows.append(
            SheetRow(
                index=row_index,
                label=label or f"Строка {row_index}",
                group=group or None,
                detail=detail or None,
                fact=fact,
                plan=plan,
                deviation=deviation,
            )
        )

        for column_index, day_offset in day_columns:
            raw_value = values_sheet.cell(row=row_index, column=column_index).value
            amount = _to_decimal(raw_value)
            if amount is None or amount == 0:
                # Нули не создают тысячи нулевых транзакций (FR-65).
                continue
            formula_value = formula_sheet.cell(row=row_index, column=column_index).value
            formula = (
                str(formula_value)
                if isinstance(formula_value, str) and formula_value.startswith("=")
                else None
            )
            occurred = start + dt.timedelta(days=day_offset)
            if occurred >= end_exclusive:
                continue
            cells.append(
                SheetDailyCell(
                    sheet_name=name,
                    row_label=label or f"Строка {row_index}",
                    row_index=row_index,
                    column_letter=_column_letter(column_index),
                    day_offset=day_offset,
                    occurred_date=occurred,
                    amount=amount,
                    formula=formula,
                )
            )

    for row_index in range(1, header_row):
        label = str(values_sheet.cell(row=row_index, column=3).value or "").strip()
        amount = _to_decimal(values_sheet.cell(row=row_index, column=6).value)
        raw_date = values_sheet.cell(row=row_index, column=5).value
        if not label or amount is None or amount == 0:
            continue
        income_date: dt.date | None = None
        if isinstance(raw_date, dt.datetime):
            income_date = raw_date.date()
        elif isinstance(raw_date, dt.date):
            income_date = raw_date
        # Недатированный доход остаётся агрегатом периода (FR-64, A54).
        incomes.append(SheetIncome(label=label, amount=amount, occurred_date=income_date))

    return ParsedSheet(
        name=name,
        period_start=start,
        period_end_exclusive=end_exclusive,
        rows=tuple(rows),
        daily_cells=tuple(cells),
        incomes=tuple(incomes),
    )


def _find_day_header_row(sheet: Any) -> int | None:
    """Строка заголовка дневной матрицы: подряд идущие номера дней."""
    for row_index in range(1, min(sheet.max_row, 40) + 1):
        numbers = 0
        for column_index in range(9, min(sheet.max_column, 60) + 1):
            value = sheet.cell(row=row_index, column=column_index).value
            if isinstance(value, int | float) and 1 <= float(value) <= 31:
                numbers += 1
        if numbers >= 20:
            return row_index
    return None


def _day_columns(sheet: Any, header_row: int) -> list[tuple[int, int]]:
    """Соответствие столбца и смещения дня от начала периода."""
    result: list[tuple[int, int]] = []
    offset = 0
    for column_index in range(9, sheet.max_column + 1):
        value = sheet.cell(row=header_row, column=column_index).value
        if not isinstance(value, int | float):
            continue
        day = int(value)
        if not 1 <= day <= 31:
            continue
        result.append((column_index, offset))
        offset += 1
    return result


def _looks_like_total_row(label: str) -> bool:
    lowered = label.casefold()
    return any(
        token in lowered
        for token in ("итого", "итог", "всего", "сумма за", "недельный", "дневной итог")
    )


def _column_letter(index: int) -> str:
    from openpyxl.utils import get_column_letter

    letter: str = get_column_letter(index)
    return letter
