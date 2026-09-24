"""Выгрузка журнала в XLSX и CSV (FR-67, SEC-08, A108, A200, NFR-12).

Экспорт доступен независимо от AI. Внешний текст, похожий на формулу,
экспортируется как текст для защиты от CSV injection.
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import uuid
from dataclasses import dataclass

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from fintracker.application.conversation.context import category_paths
from fintracker.core.money import Money
from fintracker.db.models.access import Beneficiary, Person, User, Workspace
from fintracker.db.models.catalog import Category, Tag, TransactionTag
from fintracker.db.models.commitments import Goal
from fintracker.db.models.ledger import Allocation, Transaction, TransactionRevision
from fintracker.db.models.planning import BudgetLine, BudgetPeriod, BudgetVersion, PeriodPolicyRow

# Символы, с которых начинается формула в табличных редакторах.
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def sanitize_cell(value: str | None) -> str:
    """Обезвредить внешний текст, похожий на формулу (OWASP CSV injection)."""
    if value is None:
        return ""
    text = str(value)
    if text.startswith(_FORMULA_PREFIXES):
        return "'" + text
    return text


@dataclass(frozen=True, slots=True)
class ExportRow:
    transaction_id: uuid.UUID
    revision: int
    transaction_type: str
    occurred_date: dt.date
    date_precision: str
    granularity: str
    amount_minor: int
    currency: str
    status: str
    origin: str
    actor_name: str
    spender_name: str
    note: str
    tags: str
    allocations: tuple[dict[str, str], ...]


@dataclass(frozen=True, slots=True)
class ExportSnapshot:
    workspace_name: str
    currency: str
    data_revision: int
    generated_at: dt.datetime
    rows: tuple[ExportRow, ...]
    periods: tuple[dict[str, str], ...]
    categories: tuple[dict[str, str], ...] = ()
    plans: tuple[dict[str, str], ...] = ()
    goals: tuple[dict[str, str], ...] = ()


async def build_snapshot(
    session: AsyncSession,
    *,
    workspace: Workspace,
    date_from: dt.date | None = None,
    date_to_exclusive: dt.date | None = None,
) -> ExportSnapshot:
    """Один зафиксированный снимок журнала (ADR-09)."""
    statement = (
        select(Transaction, TransactionRevision)
        .join(
            TransactionRevision,
            (TransactionRevision.workspace_id == Transaction.workspace_id)
            & (TransactionRevision.transaction_id == Transaction.id)
            & (TransactionRevision.revision == Transaction.current_revision),
        )
        .where(Transaction.workspace_id == workspace.id)
        .order_by(TransactionRevision.occurred_date, Transaction.id)
    )
    if date_from is not None:
        statement = statement.where(TransactionRevision.occurred_date >= date_from)
    if date_to_exclusive is not None:
        statement = statement.where(TransactionRevision.occurred_date < date_to_exclusive)
    rows = (await session.execute(statement)).all()

    paths = await category_paths(session, workspace_id=workspace.id)
    beneficiaries = {
        row[0]: row[1]
        for row in (
            await session.execute(
                select(Beneficiary.id, Beneficiary.name).where(
                    Beneficiary.workspace_id == workspace.id
                )
            )
        ).all()
    }
    people = {
        row[0]: row[1]
        for row in (
            await session.execute(
                select(Person.id, Person.name).where(Person.workspace_id == workspace.id)
            )
        ).all()
    }
    users = {
        row[0]: str(row[1])
        for row in (await session.execute(select(User.id, User.telegram_user_id))).all()
    }

    export_rows: list[ExportRow] = []
    for transaction, revision in rows:
        allocations = (
            (
                await session.execute(
                    select(Allocation).where(
                        Allocation.workspace_id == workspace.id,
                        Allocation.transaction_id == transaction.id,
                        Allocation.revision == revision.revision,
                    )
                )
            )
            .scalars()
            .all()
        )
        tag_names = (
            (
                await session.execute(
                    select(Tag.name)
                    .join(
                        TransactionTag,
                        (TransactionTag.workspace_id == Tag.workspace_id)
                        & (TransactionTag.tag_id == Tag.id),
                    )
                    .where(
                        TransactionTag.workspace_id == workspace.id,
                        TransactionTag.transaction_id == transaction.id,
                        TransactionTag.revision == revision.revision,
                    )
                )
            )
            .scalars()
            .all()
        )
        export_rows.append(
            ExportRow(
                transaction_id=transaction.id,
                revision=revision.revision,
                transaction_type=revision.transaction_type,
                occurred_date=revision.occurred_date,
                date_precision=revision.date_precision,
                granularity=revision.granularity,
                amount_minor=revision.amount_minor,
                currency=revision.currency,
                status=transaction.status,
                origin=transaction.origin,
                actor_name=users.get(transaction.created_by, ""),
                spender_name=people.get(revision.spender_person_id, "")
                if revision.spender_person_id
                else "",
                note=revision.note or "",
                tags="; ".join(str(name) for name in tag_names),
                allocations=tuple(
                    {
                        "allocation_id": str(item.id),
                        "stable_line_id": str(item.stable_line_id),
                        "role": item.economic_role,
                        "category": paths.get(item.category_id, "") if item.category_id else "",
                        "beneficiary": beneficiaries.get(item.beneficiary_id, "")
                        if item.beneficiary_id
                        else "",
                        "amount_minor": str(item.amount_minor),
                        "label": item.line_label or "",
                    }
                    for item in allocations
                ),
            )
        )

    periods = (
        await session.execute(
            select(BudgetPeriod, PeriodPolicyRow)
            .join(
                PeriodPolicyRow,
                (PeriodPolicyRow.workspace_id == BudgetPeriod.workspace_id)
                & (PeriodPolicyRow.id == BudgetPeriod.policy_id),
            )
            .where(BudgetPeriod.workspace_id == workspace.id)
            .order_by(BudgetPeriod.start_date)
        )
    ).all()
    categories = (
        await session.scalars(
            select(Category)
            .where(Category.workspace_id == workspace.id)
            .order_by(Category.sort_order, Category.id)
        )
    ).all()
    plans = (
        await session.execute(
            select(BudgetVersion, BudgetLine)
            .outerjoin(
                BudgetLine,
                (BudgetLine.workspace_id == BudgetVersion.workspace_id)
                & (BudgetLine.budget_version_id == BudgetVersion.id),
            )
            .where(BudgetVersion.workspace_id == workspace.id)
            .order_by(
                BudgetVersion.period_id, BudgetVersion.kind, BudgetVersion.version, BudgetLine.id
            )
        )
    ).all()
    goals = (
        await session.scalars(
            select(Goal).where(Goal.workspace_id == workspace.id).order_by(Goal.created_at, Goal.id)
        )
    ).all()
    return ExportSnapshot(
        workspace_name=workspace.name,
        currency=workspace.currency,
        data_revision=workspace.data_revision,
        generated_at=dt.datetime.now(dt.UTC),
        rows=tuple(export_rows),
        categories=tuple(
            {
                "category_id": str(row.id),
                "name": row.name,
                "path": paths.get(row.id, row.name),
                "parent_id": str(row.parent_id or ""),
                "archived_at": row.archived_at.isoformat() if row.archived_at else "",
                "merged_into_id": str(row.merged_into_id or ""),
            }
            for row in categories
        ),
        plans=tuple(
            {
                "period_id": str(version.period_id),
                "version_id": str(version.id),
                "kind": version.kind,
                "version": str(version.version),
                "status": version.plan_status,
                "origin": version.origin,
                "overall_limit_minor": str(version.overall_limit_minor)
                if version.overall_limit_minor is not None
                else "",
                "category_id": str(line.category_id) if line else "",
                "category": paths.get(line.category_id, "") if line else "",
                "beneficiary_id": str(line.beneficiary_id or "") if line else "",
                "limit_minor": str(line.limit_minor)
                if line and line.limit_minor is not None
                else "",
                "rollover_mode": line.rollover_mode if line else "",
            }
            for version, line in plans
        ),
        goals=tuple(
            {
                "goal_id": str(row.id),
                "name": row.name,
                "kind": row.kind,
                "currency": row.currency,
                "status": row.status,
                "target_minor": str(row.target_minor) if row.target_minor is not None else "",
                "allocated_minor": str(row.allocated_minor),
                "contribution_minor": str(row.contribution_minor)
                if row.contribution_minor is not None
                else "",
                "due_date": row.due_date.isoformat() if row.due_date else "",
            }
            for row in goals
        ),
        periods=tuple(
            {
                "budget_id": str(workspace.id),
                "period_id": str(period.id),
                "start_date": period.start_date.isoformat(),
                # Включённая дата конца с явным названием поля (FR-67).
                "end_date_inclusive": (period.end_exclusive - dt.timedelta(days=1)).isoformat(),
                "repeat_rule": f"{policy.mode}:{policy.interval}",
                "policy_version": str(policy.version),
                "plan_origin": period.state,
            }
            for period, policy in periods
        ),
    )


CSV_COLUMNS = (
    "transaction_id",
    "revision",
    "type",
    "date",
    "date_precision",
    "granularity",
    "amount_minor",
    "currency",
    "status",
    "origin",
    "actor",
    "spender",
    "note",
    "tags",
    "allocations_count",
)

ALLOCATION_COLUMNS = (
    "transaction_id",
    "revision",
    "allocation_id",
    "stable_line_id",
    "role",
    "category",
    "beneficiary",
    "amount_minor",
    "label",
)


def to_csv(snapshot: ExportSnapshot) -> bytes:
    """Нормализованный журнал в UTF-8 с экранированием (FR-67)."""
    buffer = io.StringIO()
    writer = csv.writer(buffer, quoting=csv.QUOTE_ALL, lineterminator="\n")
    writer.writerow(CSV_COLUMNS)
    for row in snapshot.rows:
        writer.writerow(
            [
                str(row.transaction_id),
                str(row.revision),
                row.transaction_type,
                row.occurred_date.isoformat(),
                row.date_precision,
                row.granularity,
                str(row.amount_minor),
                row.currency,
                row.status,
                row.origin,
                sanitize_cell(row.actor_name),
                sanitize_cell(row.spender_name),
                sanitize_cell(row.note),
                sanitize_cell(row.tags),
                str(len(row.allocations)),
            ]
        )
    return buffer.getvalue().encode("utf-8-sig")


def to_xlsx(snapshot: ExportSnapshot) -> bytes:
    """XLSX: журнал, распределения, периоды, категории, версии планов и цели."""
    from openpyxl import Workbook

    workbook = Workbook()
    operations = workbook.active
    operations.title = "Операции"
    operations.append(list(CSV_COLUMNS))
    for row in snapshot.rows:
        operations.append(
            [
                str(row.transaction_id),
                row.revision,
                row.transaction_type,
                row.occurred_date.isoformat(),
                row.date_precision,
                row.granularity,
                row.amount_minor,
                row.currency,
                row.status,
                row.origin,
                sanitize_cell(row.actor_name),
                sanitize_cell(row.spender_name),
                sanitize_cell(row.note),
                sanitize_cell(row.tags),
                len(row.allocations),
            ]
        )

    allocations = workbook.create_sheet("Распределения")
    allocations.append(list(ALLOCATION_COLUMNS))
    for row in snapshot.rows:
        for allocation in row.allocations:
            allocations.append(
                [
                    str(row.transaction_id),
                    row.revision,
                    allocation["allocation_id"],
                    allocation["stable_line_id"],
                    allocation["role"],
                    sanitize_cell(allocation["category"]),
                    sanitize_cell(allocation["beneficiary"]),
                    int(allocation["amount_minor"]),
                    sanitize_cell(allocation["label"]),
                ]
            )

    periods = workbook.create_sheet("Периоды")
    if snapshot.periods:
        periods.append(list(snapshot.periods[0].keys()))
        for period in snapshot.periods:
            periods.append([sanitize_cell(value) for value in period.values()])

    for title, records, empty_columns in (
        (
            "Категории",
            snapshot.categories,
            ("category_id", "name", "path", "parent_id", "archived_at", "merged_into_id"),
        ),
        (
            "Бюджеты",
            snapshot.plans,
            (
                "period_id",
                "version_id",
                "kind",
                "version",
                "status",
                "origin",
                "overall_limit_minor",
                "category_id",
                "category",
                "beneficiary_id",
                "limit_minor",
                "rollover_mode",
            ),
        ),
        (
            "Цели",
            snapshot.goals,
            (
                "goal_id",
                "name",
                "kind",
                "currency",
                "status",
                "target_minor",
                "allocated_minor",
                "contribution_minor",
                "due_date",
            ),
        ),
    ):
        sheet = workbook.create_sheet(title)
        sheet.append(list(records[0]) if records else list(empty_columns))
        for record in records:
            sheet.append([sanitize_cell(value) for value in record.values()])

    description = workbook.create_sheet("Описание полей")
    description.append(["Поле", "Значение"])
    for field_name, explanation in (
        ("amount_minor", "Сумма в минимальных единицах валюты, целое число"),
        ("granularity", "individual — отдельная операция; daily_aggregate — дневной итог"),
        ("date_precision", "Точность даты: day, time, interval, unknown"),
        ("status", "posted — учитывается; voided — отменена"),
        ("role", "Экономическая роль части: expense, expense_refund и другие"),
        (
            "Операции и Распределения",
            "Суммы распределений не складываются с суммой операции: "
            "это разные листы одного события",
        ),
        (
            "Бюджеты",
            "Все версии планов: baseline — исходный, working — рабочий. "
            "Для текущих лимитов используйте последнюю рабочую версию периода; "
            "версии не складываются.",
        ),
        ("limit_minor", "Пусто — лимит не задан; 0 — расходы не запланированы"),
        ("allocated_minor", "Резерв цели, не подтверждённый баланс банковского счёта"),
        ("end_date_inclusive", "Дата конца периода включительно"),
        ("Валюта", snapshot.currency),
        ("Версия данных", str(snapshot.data_revision)),
        ("Снимок", snapshot.generated_at.isoformat()),
    ):
        description.append([field_name, sanitize_cell(explanation)])

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def build_matrix(
    snapshot: ExportSnapshot, *, start: dt.date, end_inclusive: dt.date
) -> list[list[str]]:
    """Привычная матрица «категории × дни» по реальным датам интервала (FR-67).

    Ограничения в 31 колонку нет: длинный отчёт сохраняет все дни.
    """
    days = [start + dt.timedelta(days=offset) for offset in range((end_inclusive - start).days + 1)]
    categories: dict[str, dict[dt.date, int]] = {}
    for row in snapshot.rows:
        if row.status != "posted":
            continue
        if not (start <= row.occurred_date <= end_inclusive):
            continue
        for allocation in row.allocations:
            if allocation["role"] not in {"expense", "expense_refund"}:
                continue
            name = allocation["category"] or "Без категории"
            signed = int(allocation["amount_minor"])
            if allocation["role"] == "expense_refund":
                signed = -signed
            bucket = categories.setdefault(name, {})
            bucket[row.occurred_date] = bucket.get(row.occurred_date, 0) + signed

    header = ["Категория", *[day.isoformat() for day in days], "Итого"]
    matrix = [header]
    for name in sorted(categories):
        values = categories[name]
        line = [sanitize_cell(name)]
        total = 0
        for day in days:
            amount = values.get(day, 0)
            total += amount
            line.append(Money(amount, snapshot.currency).format(with_symbol=False))
        line.append(Money(total, snapshot.currency).format(with_symbol=False))
        matrix.append(line)
    return matrix
