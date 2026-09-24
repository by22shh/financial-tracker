"""Telegram audit: faithful report/export and confirmed category rollovers."""

from __future__ import annotations

import datetime as dt
import io
from dataclasses import replace

import pytest
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from fintracker.application.analytics.reports import _label_for, spending_report
from fintracker.application.commitments.goals import create_goal
from fintracker.application.conversation.keyboards import short
from fintracker.application.conversation.rollover_flow import rollover_action
from fintracker.application.integrations.exporter import build_snapshot, to_xlsx
from fintracker.application.ledger.service import post_transaction
from fintracker.application.planning.periods import period_for_date
from fintracker.application.planning.plan import PlanLineSpec, create_budget_version, period_status
from fintracker.application.planning.rollover_actions import accept_rollover
from fintracker.config import Settings
from fintracker.core.errors import ConflictError, ValidationFailed
from fintracker.core.money import Money
from fintracker.db.models.planning import Rollover
from tests.conftest import requires_pg
from tests.integration.factories import build_fixture
from tests.integration.test_money_scenarios import expense_spec

pytestmark = [pytest.mark.pg, requires_pg]


async def test_uncategorized_amount_survives_report_grouping(owner_session: AsyncSession) -> None:
    fixture = await build_fixture(owner_session)
    categorized = expense_spec(fixture, amount=Money(90000, "RUB"), category="Продукты")
    uncategorized = expense_spec(fixture, amount=Money(115000, "RUB"), category="Продукты")
    uncategorized = replace(
        uncategorized, allocations=(replace(uncategorized.allocations[0], category_id=None),)
    )
    for spec in (categorized, uncategorized):
        await post_transaction(
            owner_session, fixture.uow, actor=fixture.actor, spec=spec, origin="form"
        )
    for grouping in ("category", "none", "beneficiary"):
        report = await spending_report(
            owner_session,
            workspace=fixture.workspace,
            date_from=fixture.period.start_date,
            date_to_exclusive=fixture.period.end_exclusive,
            group_by=grouping,
        )
        assert report.total_minor == 205000
        assert report.uncategorized_minor == 115000
        if grouping == "category":
            assert {row.label for row in report.rows} == {"Без категории", "Продукты"}
        if grouping == "none":
            assert report.rows[0].label == "Всего"


async def test_export_includes_real_categories_plan_and_goal(owner_session: AsyncSession) -> None:
    fixture = await build_fixture(owner_session, limits={"Продукты": 12345})
    await create_goal(
        owner_session,
        fixture.uow,
        actor=fixture.actor,
        name="=unsafe goal",
        currency="RUB",
        target=Money(50000, "RUB"),
    )
    snapshot = await build_snapshot(owner_session, workspace=fixture.workspace)
    workbook = load_workbook(io.BytesIO(to_xlsx(snapshot)), data_only=False)
    assert {"Категории", "Бюджеты", "Цели"} <= set(workbook.sheetnames)
    categories = list(workbook["Категории"].values)
    assert "Продукты" in {row[1] for row in categories[1:]}
    plans = list(workbook["Бюджеты"].values)
    assert "12345" in {row[10] for row in plans[1:]}
    goals = list(workbook["Цели"].values)
    assert goals[1][1] == "'=unsafe goal"
    assert goals[1][5] == "50000"
    assert _label_for((None, None), {"categories": {}, "beneficiaries": {}}) == "Без категории"
    assert (
        _label_for((None, None), {"categories": {}, "beneficiaries": {}}, grouping="none")
        == "Всего"
    )


async def test_rollover_revalidates_and_accepts_once(owner_session: AsyncSession) -> None:
    fixture = await build_fixture(owner_session, limits={"Продукты": 10000})
    next_period = await period_for_date(
        owner_session, workspace_id=fixture.workspace.id, day=fixture.period.end_exclusive
    )
    spec = PlanLineSpec(category_id=fixture.categories["Продукты"], limit_minor=20000)
    await create_budget_version(
        owner_session,
        workspace_id=fixture.workspace.id,
        period_id=next_period.id,
        kind="working",
        plan_status="approved",
        origin="manual",
        lines=[spec],
    )
    proposal = Rollover(
        workspace_id=fixture.workspace.id,
        source_period_id=fixture.period.id,
        destination_period_id=next_period.id,
        stable_line_id=spec.resolved_stable_id(fixture.workspace.id),
        amount_minor=10000,
        mode="positive_only",
        status="proposed",
        basis_completeness="unknown",
    )
    owner_session.add(proposal)
    await owner_session.flush()
    with pytest.raises(ValidationFailed, match="после окончания"):
        await accept_rollover(
            owner_session,
            fixture.uow,
            actor=fixture.actor,
            rollover_id=proposal.id,
            today=fixture.period.start_date,
            expected_amount_minor=10000,
        )
    with pytest.raises(ConflictError, match="изменилась"):
        await accept_rollover(
            owner_session,
            fixture.uow,
            actor=fixture.actor,
            rollover_id=proposal.id,
            today=fixture.period.end_exclusive,
            expected_amount_minor=9000,
        )
    for _ in range(2):
        await accept_rollover(
            owner_session,
            fixture.uow,
            actor=fixture.actor,
            rollover_id=proposal.id,
            today=fixture.period.end_exclusive,
            expected_amount_minor=10000,
        )
    status = await period_status(
        owner_session,
        workspace_id=fixture.workspace.id,
        period_id=next_period.id,
        currency="RUB",
        today=fixture.period.end_exclusive,
    )
    assert status.lines[0].effective_limit_minor == 30000
    assert len((await owner_session.scalars(select(Rollover))).all()) == 1


async def test_closed_period_rollover_flow_rejects_changed_source(
    owner_session: AsyncSession,
    test_settings: Settings,
) -> None:
    fixture = await build_fixture(
        owner_session, start=dt.date(2026, 8, 10), limits={"Продукты": 10000}
    )
    destination = await period_for_date(
        owner_session, workspace_id=fixture.workspace.id, day=fixture.period.end_exclusive
    )
    spec = PlanLineSpec(category_id=fixture.categories["Продукты"], limit_minor=20000)
    await create_budget_version(
        owner_session,
        workspace_id=fixture.workspace.id,
        period_id=destination.id,
        kind="working",
        plan_status="approved",
        origin="manual",
        lines=[spec],
    )
    await owner_session.commit()
    replies = await rollover_action(
        test_settings,
        actor=fixture.actor,
        workspace=fixture.workspace,
        action="pick",
        rest=[short(fixture.period.id), short(spec.resolved_stable_id(fixture.workspace.id))],
    )
    accept_data = replies[0].buttons[0][0].data.split(":")
    assert accept_data[:2] == ["roll", "accept"]
    assert "100" in replies[0].text
    await post_transaction(
        owner_session,
        fixture.uow,
        actor=fixture.actor,
        origin="form",
        spec=expense_spec(
            fixture,
            amount=Money(1000, "RUB"),
            category="Продукты",
            occurred=fixture.period.start_date,
        ),
    )
    await owner_session.commit()
    with pytest.raises(ConflictError, match="Остаток изменился"):
        await rollover_action(
            test_settings,
            actor=fixture.actor,
            workspace=fixture.workspace,
            action="accept",
            rest=accept_data[2:],
        )
    replies = await rollover_action(
        test_settings,
        actor=fixture.actor,
        workspace=fixture.workspace,
        action="pick",
        rest=[short(fixture.period.id), short(spec.resolved_stable_id(fixture.workspace.id))],
    )
    new_data = replies[0].buttons[0][0].data.split(":")
    with pytest.raises(ConflictError, match="Сумма предложения изменилась"):
        await rollover_action(
            test_settings,
            actor=fixture.actor,
            workspace=fixture.workspace,
            action="accept",
            rest=accept_data[2:],
        )
    for _ in range(2):
        replies = await rollover_action(
            test_settings,
            actor=fixture.actor,
            workspace=fixture.workspace,
            action="accept",
            rest=new_data[2:],
        )
        assert "Перенос принят" in replies[0].text
    owner_session.expire_all()
    status = await period_status(
        owner_session,
        workspace_id=fixture.actor.require_workspace(),
        period_id=destination.id,
        currency="RUB",
        today=dt.date(2026, 9, 20),
    )
    assert status.lines[0].rollover_minor == 9000
    assert status.lines[0].effective_limit_minor == 29000
